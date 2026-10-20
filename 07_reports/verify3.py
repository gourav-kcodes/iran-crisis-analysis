# Independent re-computation of key report numbers from the raw CSVs (does not use analysis_core.py), checked
# against the text of the three analysis reports built by build_reports.py.
# Usage (project root): python 07_reports/verify3.py <best.docx> <zs.docx> <fs.docx>
import re, sys, numpy as np, pandas as pd
from docx import Document


def doc_text(path):
    d = Document(path)
    t = '\n'.join(p.text for p in d.paragraphs)
    for tb in d.tables:
        for r in tb.rows:
            t += '\n' + ' | '.join(c.text for c in r.cells)
    return t


TEXT = dict(zip(['best', 'zs', 'fs'], (doc_text(p) for p in sys.argv[1:4])))
from scipy.stats import fisher_exact, mannwhitneyu, norm
from sklearn.metrics import cohen_kappa_score, f1_score
D='data/clean/'
def L(n):
    d=pd.read_csv(D+f'iran_crisis_videos_classified_{n}.csv',low_memory=False,dtype={'id':str})
    d['s']=d['sentiment_label'].astype(str).str.strip().str.replace('Humor','Humour'); d['f']=d['frame_label'].astype(str).str.strip().str.replace('Humor','Humour')
    return d
mz,mf,qz,qf=L('mistral'),L('mistral_fewshot'),L('qwen'),L('qwen_fewshot')
g=pd.read_csv(D+'gold_standard_dataset.csv'); g['k']=g['id'].astype(str).str.replace(r'\D','',regex=True).str[:18]
fails=0
def chk(S,txt,label):
    global fails
    T=TEXT[S]
    ok=txt in T
    if not ok: fails+=1
    print(('OK  ' if ok else 'MISS'),S,label,'->',txt)
for S in ['best','zs','fs']:
    if S=='best': M=mz.merge(mf[['id','f']].rename(columns={'f':'f2'}),on='id'); M['f']=M['f2']; Q=qf
    elif S=='zs': M,Q=mz,qz
    else: M,Q=mf,qf
    n=len(M)
    chk(S,f"labels only {100*(M.s=='Neutral').mean():.1f}% Neutral",'neutral share')
    chk(S,f"{100*(M.f=='Political blame').mean():.1f}%",'PB share')
    # gold kappa
    k=M.assign(k=M['id'].str.replace(r'\D','',regex=True).str[:18]).drop_duplicates('k')
    gm=g.merge(k[['k','s','f']],on='k')
    gs=gm['tier1_sentiment'].str.strip().str.replace('Humor','Humour'); gf=gm['tier3_frame'].str.strip()
    chk(S,f"| {cohen_kappa_score(gs,gm.s):.3f} | {f1_score(gs,gm.s,average='weighted',zero_division=0):.3f} | {f1_score(gs,gm.s,average='macro',zero_division=0):.3f}",'T1 kappa/wf1/mf1')
    chk(S,f"| {cohen_kappa_score(gf,gm.f):.3f} | {f1_score(gf,gm.f,average='weighted',zero_division=0):.3f}",'T3 kappa/wf1')
    chk(S,f"Political blame to {100*(gm.f=='Political blame').mean():.1f}%",'PB on gold')
    # Anger -> PB
    a=M[M.s=='Anger']; v=100*(a.f=='Political blame').mean(); lift=v/(100*(M.f=='Political blame').mean())
    chk(S,f"{v:.1f}%",'anger->pb'); chk(S,f"lift {lift:.1f}",'anger->pb lift')
    # Indonesian vs English religious
    i=M[M.language_full=='Indonesian']; e=M[M.language_full=='English']
    ai,ae=(i.f=='Religious framing').sum(),(e.f=='Religious framing').sum()
    p=fisher_exact([[ai,len(i)-ai],[ae,len(e)-ae]])[1]
    chk(S,f"{100*ai/len(i):.1f}% ({ai}/{len(i)})",'ID rel'); print('   fisher p',p)
    # GB vs US solidarity
    gb=M[M.region_code=='GB']; us=M[M.region_code=='US']
    chk(S,f"{100*(gb.s=='Solidarity').mean():.1f}% vs {100*(us.s=='Solidarity').mean():.1f}%",'GB/US sol')
    # trend test solidarity
    d=M.copy(); d['dt']=pd.to_datetime(d['create_time_readable'],errors='coerce'); d=d.dropna(subset=['dt']); d['w']=d.dt.dt.to_period('W')
    c=d.w.value_counts(); d=d[d.w.isin(c[c>=30].index)]
    ws=sorted(d.w.unique()); t=np.arange(len(ws)); nn=np.array([(d.w==w).sum() for w in ws]); x=np.array([((d.w==w)&(d.s=='Solidarity')).sum() for w in ws])
    pbar=x.sum()/nn.sum(); T=(t*x).sum()-pbar*(t*nn).sum(); V=pbar*(1-pbar)*((t*t*nn).sum()-(t*nn).sum()**2/nn.sum())
    chk(S,f"Solidarity | {100*x[0]/nn[0]:.1f}% | {100*x[-1]/nn[-1]:.1f}% | {T/np.sqrt(V):.2f}",'sol trend z')
    # engagement
    sol=M[M.s=='Solidarity'].like_count; rest=M[(M.s!='Solidarity')&(M.s!='Unclear')].like_count
    chk(S,f"highest median likes ({np.median(sol):g} [",'sol likes'); print('   MW p',mannwhitneyu(sol,rest).pvalue)
    ec=M[M.f=='Economic/sanctions'].like_count
    chk(S,f"lowest median likes ({np.median(ec):g}, p < 0.001)",'econ likes')
    # weekly fear peak share
    fw=d.groupby('w').apply(lambda z:100*(z.s=='Fear').mean(),include_groups=False)
    chk(S,f"({fw.iloc[5]:.1f}%)",'fear wk6')
    # non-English gold unclear
    ne=gm[gm.language_full!='English']
    chk(S,f"{100*(ne.tier3_frame.str.strip()=='Unclear').mean():.1f}% of the {len(ne)} non-English gold rows",'ne gold unclear')
print('FAILS',fails)
# independent check of the classification-error correction (plain crosstab implementation, no bootstrap)
for S in ['best','zs','fs']:
    if S=='best': M=mz.merge(mf[['id','f']].rename(columns={'f':'f2'}),on='id'); M['f']=M['f2']
    elif S=='zs': M=mz
    else: M=mf
    k=M.assign(k=M['id'].str.replace(r'\D','',regex=True).str[:18]).drop_duplicates('k')
    gm=g.merge(k[['k','s','f']],on='k'); gm['tf']=gm['tier3_frame'].str.strip()
    rate=pd.crosstab(gm.f,gm.tf,normalize='index')
    def corr(sub,lab):
        q=sub.f.value_counts(normalize=True)
        return 100*sum(q.get(j,0)*(rate.loc[j,lab] if lab in rate.columns else 0) for j in rate.index)
    T=TEXT[S]
    for txt in [f"Political blame falls from {100*(M.f=='Political blame').mean():.1f}% to {corr(M,'Political blame'):.1f}%",
                f"{corr(M[M.language_full=='Indonesian'],'Religious framing'):.1f}% vs {corr(M[M.language_full=='English'],'Religious framing'):.1f}%"]:
        ok=txt in T; fails+=0 if ok else 1
        print('OK  ' if ok else 'MISS', S, 'calib ->', txt)
print('FAILS',fails)
