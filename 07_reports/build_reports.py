# Run analysis_core.py and report_figs.py for the setting first.
# Usage (project root): python 07_reports/build_reports.py best|zs|fs <template.docx> <out.docx>

import sys, copy
from docx import Document
from docx.shared import Inches
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from xml.sax.saxutils import escape

S, TEMPLATE, OUT = sys.argv[1:4]
FIG = f'data/analysis/report_{S}'
doc = Document(TEMPLATE)
body = doc.element.body
for el in list(body):
    if el.tag != qn('w:sectPr'):
        body.remove(el)
SECT = body.find(qn('w:sectPr'))
FONT = '<w:rFonts w:ascii="Calibri" w:cs="Calibri" w:eastAsia="Calibri" w:hAnsi="Calibri"/>'


def add(xml):
    el = parse_xml(xml)
    SECT.addprevious(el)
    return el


def run(text, bold=False, italic=False, size=22, color='000000'):
    b = '<w:b/><w:bCs/>' if bold else '<w:b w:val="false"/><w:bCs w:val="false"/>'
    i = '<w:i/><w:iCs/>' if italic else '<w:i w:val="false"/><w:iCs w:val="false"/>'
    return (f'<w:r><w:rPr>{FONT}{b}{i}<w:color w:val="{color}"/><w:sz w:val="{size}"/><w:szCs w:val="{size}"/></w:rPr>'
            f'<w:t xml:space="preserve">{escape(text)}</w:t></w:r>')


def spacer(before=60, keep=False):
    k = '<w:keepNext/>' if keep else ''
    add(f'<w:p {nsdecls("w")}><w:pPr>{k}<w:spacing w:after="0" w:before="{before}"/></w:pPr></w:p>')


def title(t, sub):
    add(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:after="60" w:before="0"/><w:jc w:val="center"/></w:pPr>{run(t, True, size=44)}</w:p>')
    add(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:after="200" w:before="0"/><w:jc w:val="center"/></w:pPr>'
        f'{run(sub, italic=True, color="555555")}</w:p>')


def h1(t):
    add(f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="Heading1"/><w:keepNext/><w:spacing w:after="120" w:before="300"/></w:pPr>{run(t, True, size=28)}</w:p>')


def h2(t):
    spacer(100, keep=True)
    add(f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="Heading2"/><w:keepNext/><w:spacing w:after="80" w:before="200"/></w:pPr>{run(t, True, size=24)}</w:p>')


def para(text, lead=None, italic=False):
    r = (run(lead, True) if lead else '') + run(text.strip(), italic=italic)
    add(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:after="60" w:before="60"/></w:pPr>{r}</w:p>')


def bullet(text, lead=None):
    r = (run(lead, True) if lead else '') + run(text)
    add(f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="ListParagraph"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="2"/></w:numPr>'
        f'<w:spacing w:after="60" w:before="60"/></w:pPr>{r}</w:p>')


FN = [0]


def fig(n, caption, width=6.04):
    spacer(100, keep=True)
    p = add(f'<w:p {nsdecls("w")}><w:pPr><w:keepNext/><w:spacing w:after="60" w:before="100"/><w:jc w:val="center"/></w:pPr></w:p>')
    from docx.text.paragraph import Paragraph
    Paragraph(p, doc._body).add_run().add_picture(f'{FIG}/fig{n:02d}.png', width=Inches(width))
    FN[0] += 1
    assert FN[0] == n, (FN[0], n)
    add(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:after="140" w:before="0"/><w:jc w:val="center"/></w:pPr>'
        f'{run(f"Figure {n}. {caption}", italic=True, size=19, color="555555")}</w:p>')


def cell(text, w, header=False, align='center', size=22):
    fill = 'F2F2F2' if header else 'FFFFFF'
    m = 80 if header else 70
    return (f'<w:tc><w:tcPr><w:tcW w:type="dxa" w:w="{w}"/><w:tcBorders>'
            + ''.join(f'<w:{s} w:val="single" w:color="000000" w:sz="4"/>' for s in ('top', 'left', 'bottom', 'right'))
            + f'</w:tcBorders><w:shd w:fill="{fill}" w:val="clear"/><w:tcMar><w:top w:type="dxa" w:w="{m}"/>'
            f'<w:left w:type="dxa" w:w="120"/><w:bottom w:type="dxa" w:w="{m}"/><w:right w:type="dxa" w:w="120"/></w:tcMar></w:tcPr>'
            f'<w:p><w:pPr><w:spacing w:after="0" w:before="0"/><w:jc w:val="{align}"/></w:pPr>{run(text, bold=header, size=size)}</w:p></w:tc>')


def table(headers, rows, widths, size=22):
    spacer(80, keep=True)
    g = ''.join(f'<w:gridCol w:w="{w}"/>' for w in widths)
    hdr = '<w:tr><w:trPr><w:tblHeader/></w:trPr>' + ''.join(cell(h, w, True, 'left' if i == 0 else 'center', size)
                                                            for i, (h, w) in enumerate(zip(headers, widths))) + '</w:tr>'
    body_rows = ''.join('<w:tr>' + ''.join(cell(v, w, False, 'left' if i == 0 else 'center', size)
                                           for i, (v, w) in enumerate(zip(r, widths))) + '</w:tr>' for r in rows)
    add(f'<w:tbl {nsdecls("w")}><w:tblPr><w:tblW w:type="dxa" w:w="{sum(widths)}"/><w:tblBorders>'
        + ''.join(f'<w:{s} w:val="single" w:color="auto" w:sz="4"/>' for s in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'))
        + f'</w:tblBorders><w:jc w:val="center"/></w:tblPr><w:tblGrid>{g}</w:tblGrid>{hdr}{body_rows}</w:tbl>')


def caption(text):
    add(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:after="140" w:before="60"/><w:jc w:val="center"/></w:pPr>'
        f'{run(text, italic=True, size=19, color="555555")}</w:p>')


import report_text
report_text.build(sys.modules[__name__], S)
doc.save(OUT)
print('saved', OUT, 'figures', FN[0])
