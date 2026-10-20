# Independent check of the final report: recomputes key numbers straight from the CSVs (not from results.json)
# and checks that each appears in the report text. Usage (project root): python 07_reports/verify_final.py <Final_Report.docx>
import sys, re
import numpy as np, pandas as pd
from docx import Document
from scipy import stats
from sklearn.metrics import cohen_kappa_score as K, f1_score

doc = Document(sys.argv[1])
txt = '\n'.join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for r in t.rows:
        txt += '\n' + ' | '.join(c.text for c in r.cells)
nid = lambda s: re.sub(r'[^0-9]', '', str(s))[:18]
nl = lambda s: str(s).strip().replace('Humor', 'Humour')
D = 'data/clean/'
FAIL = []


def chk(s, why):
    ok = s in txt
    print(('OK  ' if ok else 'FAIL'), why, '->', s)
    if not ok:
        FAIL.append(why)


def load(f):
    d = pd.read_csv(D + f'iran_crisis_videos_classified_{f}.csv', low_memory=False, dtype={'id': str})
    d['k'] = d['id'].map(nid)
    for c in ('sentiment_label', 'frame_label'):
        d[c] = d[c].astype(str).map(nl)
    return d


g = pd.read_csv(D + 'gold_standard_dataset.csv')
g['k'] = g['id'].map(nid)
g['s'], g['f'] = g['tier1_sentiment'].map(nl), g['tier3_frame'].map(nl)
C = {f: load(f) for f in ('mistral', 'mistral_fewshot', 'qwen', 'qwen_fewshot', 'deepseek', 'deepseek_fewshot')}
names = {'mistral': ('Mistral 7B', 'Zero-shot'), 'mistral_fewshot': ('Mistral 7B', 'Few-shot'), 'qwen': ('Qwen2.5-3B', 'Zero-shot'),
         'qwen_fewshot': ('Qwen2.5-3B', 'Few-shot'), 'deepseek': ('DeepSeek-R1', 'Zero-shot'), 'deepseek_fewshot': ('DeepSeek-R1', 'Few-shot')}
for f, d in C.items():
    m = g.merge(d.drop_duplicates('k')[['k', 'sentiment_label', 'frame_label']], on='k')
    k1, k3 = K(m['s'], m['sentiment_label']), K(m['f'], m['frame_label'])
    f1, f3 = f1_score(m['s'], m['sentiment_label'], average='macro'), f1_score(m['f'], m['frame_label'], average='macro')
    nm, st = names[f]
    chk(f'{nm} | {st} | Tier 1 | {k1:.3f} |', f'{f} T1 kappa (App. C)')
    chk(f'{f1:.3f} | ', f'{f} T1 macro F1')
    chk(f'{nm} | {st} | Tier 3 | {k3:.3f} |', f'{f} T3 kappa (App. C)')
    chk(f'{f3:.3f}', f'{f} T3 macro F1')
# collection
M = C['mistral'].merge(C['mistral_fewshot'][['id', 'frame_label']].rename(columns={'frame_label': 'fl'}), on='id')
M['frame_label'] = M['fl']
mq = M['date_range'].isin({'20260301_20260331', '20260401_20260430'})
chk(f'({(~mq).sum():,} videos), plus one query for each whole month ({mq.sum():,} videos)', 'collection split')
dt = pd.to_datetime(M['create_time_readable'], errors='coerce')
chk(f'put {((dt.dt.date.astype(str) == "2026-03-31") & mq).sum()} videos on March 31 and '
    f'{((dt.dt.date.astype(str) == "2026-04-30") & mq).sum()} on April 30', 'month-end counts')
# corrected frame shares (calibration estimator, point estimate)
for tier, gc, pc_ in (('frame', 'f', 'frame_label'), ('sent', 's', 'sentiment_label')):
    gm = g.merge(M.drop_duplicates('k')[['k', pc_]], on='k')
    q = M[pc_].value_counts(normalize=True)
    cond = pd.crosstab(gm[pc_], gm[gc], normalize='index')
    corr = sum(q.get(j, 0) * cond.loc[j] for j in cond.index) * 100
    if tier == 'frame':
        chk(f'Military/strategic is the most common frame ({corr["Military/strategic"]:.1f}%)', 'corrected Mil/strat')
        chk(f'Political blame falls from {q["Political blame"] * 100:.1f}% to {corr["Political blame"]:.1f}%', 'corrected PB')
# weekly
M['week'] = dt.dt.to_period('W')
wk = M.dropna(subset=['week'])
cnt = wk['week'].value_counts()
weeks = sorted(cnt[cnt >= 30].index)
sh = lambda col, lab, w: (wk[wk['week'] == w][col] == lab).mean() * 100
chk(f'Solidarity is highest in the first week ({sh("sentiment_label", "Solidarity", weeks[0]):.1f}%) and falls to '
    f'{sh("sentiment_label", "Solidarity", weeks[-1]):.1f}%', 'Solidarity weekly')
chk(f'Political blame rises from {sh("frame_label", "Political blame", weeks[0]):.1f}% in the first week to '
    f'{sh("frame_label", "Political blame", weeks[5]):.1f}% in April 6–12', 'PB weekly')
# religious framing Indonesian vs English
A_, B_ = M[M['language_full'] == 'Indonesian'], M[M['language_full'] == 'English']
a, b = (A_['frame_label'] == 'Religious framing').sum(), (B_['frame_label'] == 'Religious framing').sum()
p = stats.fisher_exact([[a, len(A_) - a], [b, len(B_) - b]])[1]
chk(f'Religious framing is {a / len(A_) * 100:.1f}%', 'Indonesian RF share')
chk(f'{a / len(A_) * 100:.1f}% ({a}/{len(A_)})', 'App. E Indonesian RF count')
assert p < 0.001
# engagement
med = lambda col, lab, met: M[M[col] == lab][met].median()
chk(f'fewest likes (median {med("frame_label", "Economic/sanctions", "like_count"):g}', 'Econ median likes')
chk(f'most views (median {med("frame_label", "Military/strategic", "view_count"):g}', 'Mil median views')
chk(f'most likes (median {med("sentiment_label", "Solidarity", "like_count"):g}', 'Solidarity median likes')
chk(f'the most views (median {med("sentiment_label", "Fear", "view_count"):g}', 'Fear median views')
# non-English gold
gm = g.merge(M.drop_duplicates('k')[['k', 'frame_label']], on='k')
ne = gm[gm['language_full'] != 'English']
en = gm[gm['language_full'] == 'English']
chk(f'it codes {(ne["f"] == "Unclear").mean() * 100:.1f}% of the {len(ne)} non-English rows as Unclear, against '
    f'{(en["f"] == "Unclear").mean() * 100:.1f}% of the {len(en)} English rows', 'non-English gold Unclear')
pu = ne[ne['frame_label'] == 'Unclear']
chk(f'{(pu["f"] == "Unclear").mean() * 100:.1f}% of Mistral\'s Unclear predictions', 'Unclear precision non-English')
# gold
chk(f'{(g["s"] == "Neutral").sum()} of 400 rows ({(g["s"] == "Neutral").mean() * 100:.1f}%) are Neutral', 'gold Neutral')
kr = open('data/clean/kappa_results_full.txt').read()
assert 'Cohen\'s Kappa  : 0.8455' in kr and 'Cohen\'s Kappa  : 0.8869' in kr and '393/400 (98.2%)' in kr and '366/400 (91.5%)' in kr
assert 'Total disagreements: 7 / 400' in kr and 'Total disagreements: 34 / 400' in kr
chk('κ = 0.85 for Tier 1 (98.2% raw agreement) and κ = 0.89 for Tier 3 (91.5%)', 'coder agreement')
# languages
lv = M['language_full'].value_counts()
chk(f'{lv["English"]:,} ({lv["English"] / len(M) * 100:.1f}%) are in English', 'English count')
print('FAILS', len(FAIL))
